"""Load text content from Word (.docx) files."""
from pathlib import Path
from typing import List

from docx import Document as DocxDocument

from rag.schema import Document


class WordLoader:
    """Loads a .docx file, extracting both paragraph text and table content."""

    def load(self, file_path: str) -> List[Document]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Word document not found: {file_path}")

        docx_doc = DocxDocument(str(path))

        paragraphs = [p.text.strip() for p in docx_doc.paragraphs if p.text.strip()]

        table_lines = []
        for table in docx_doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_lines.append(" | ".join(cells))

        full_text = "\n".join(paragraphs)
        if table_lines:
            full_text += "\n\n" + "\n".join(table_lines)

        return [Document(content=full_text.strip(), metadata={"source": path.name, "type": "docx"})]
